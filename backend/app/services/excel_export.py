"""Excel workbook builders for exports and the order-email attachment.

Every workbook is built from live database rows — nothing here invents a value.
A thin styling layer keeps the files presentable in Excel/LibreOffice without
depending on anything beyond openpyxl (already in requirements.txt).
"""
from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Brand palette used on exported headers.
HEADER_FILL = PatternFill("solid", fgColor="7A1F2B")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
MUTED = Font(color="666666", size=9)
WRAP = Alignment(vertical="top", wrap_text=True)


def _style_sheet(ws, headers: list[str], widths: list[int | None],
                 freeze: str = "A2") -> None:
    """Header row + column widths for a fresh sheet."""
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        if widths and col <= len(widths) and widths[col - 1]:
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.freeze_panes = freeze


def _f(value) -> str | float | None:
    """A cell value that Excel will display without turning numbers to text."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def _sheet(workbook: Workbook, title: str) -> object:
    ws = workbook.active if workbook.worksheets else workbook.create_sheet()
    ws.title = title[:31]
    return ws


def workbook_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- inventory


def inventory_workbook(rows: list[dict]) -> bytes:
    """Current inventory positions. `rows` are build_rows() output."""
    wb = Workbook()
    ws = _sheet(wb, "Inventory")
    headers = [
        "Product Image URL", "Product Name", "SKU", "Category", "Store",
        "Current Stock", "Available Stock", "MRP", "Inventory Value",
        "Stock Status", "Low Stock Threshold", "Reorder Point", "Safety Stock",
        "Demand / Day", "Days of Cover", "Risk Level",
    ]
    widths = [44, 34, 18, 22, 24, 14, 14, 12, 16, 14, 16, 14, 14, 12, 12, 12]
    _style_sheet(ws, headers, widths)

    for row in rows:
        ws.append([
            _f(row.get("image_url")),
            _f(row.get("name")),
            _f(row.get("sku")),
            _f(row.get("category_name")),
            _f(row.get("store_name")),
            _f(row.get("current_stock")),
            _f(row.get("available_stock")),
            _f(row.get("mrp")),
            _f(row.get("inventory_value")),
            _f(row.get("stock_status")),
            _f(row.get("low_stock_threshold")),
            _f(row.get("reorder_point")),
            _f(row.get("safety_stock")),
            _f(row.get("avg_daily_demand")),
            _f(row.get("days_of_stock")),
            _f(row.get("risk_level")),
        ])
    return workbook_bytes(wb)


# ---------------------------------------------------------------- products


def products_workbook(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "Products")
    headers = [
        "SKU", "Barcode", "Product Name", "Category", "Subcategory",
        "MRP", "Cost Price", "Image URL", "Status",
        "Current Quantity", "Low Stock Threshold", "Reorder Point",
        "Safety Stock", "Target Stock", "Inventory Records",
    ]
    widths = [18, 18, 34, 22, 22, 12, 12, 44, 12, 14, 16, 14, 14, 14, 14]
    _style_sheet(ws, headers, widths)

    for row in rows:
        categories = row.get("categories") or {}
        subcategories = row.get("subcategories") or {}
        ws.append([
            _f(row.get("sku")),
            _f(row.get("barcode")),
            _f(row.get("name")),
            _f(categories.get("name")),
            _f(subcategories.get("name")),
            _f(row.get("mrp")),
            _f(row.get("cost_price")),
            _f(row.get("image_url")),
            _f(row.get("status")),
            _f(row.get("available_stock")),
            _f(row.get("low_stock_threshold")),
            _f(row.get("reorder_point")),
            _f(row.get("safety_stock")),
            _f(row.get("target_stock")),
            _f(row.get("inventory_records")),
        ])
    return workbook_bytes(wb)


# ---------------------------------------------------------------- orders


def orders_workbook(rows: list[dict]) -> bytes:
    """One row per order line item. `rows` come from the export endpoint."""
    wb = Workbook()
    ws = _sheet(wb, "Orders")
    headers = [
        "Order ID", "Product", "SKU", "Quantity", "Supplier", "Store",
        "Order Status", "Order Date", "Expected Date", "Received Date",
        "Priority", "Total Value", "Created By",
    ]
    widths = [22, 34, 18, 12, 24, 24, 18, 14, 14, 14, 12, 14, 24]
    _style_sheet(ws, headers, widths)

    for order in rows:
        items = order.get("purchase_order_items") or [{}]
        for item in items:
            product = item.get("products") or {}
            ws.append([
                _f(order.get("po_number")),
                _f(product.get("name")),
                _f(product.get("sku")),
                _f(item.get("quantity")),
                _f((order.get("suppliers") or {}).get("name")),
                _f((order.get("stores") or {}).get("name")),
                _f(order.get("status")),
                _f(order.get("requested_date")),
                _f(order.get("expected_date")),
                _f(order.get("received_date")),
                _f(order.get("priority")),
                _f(order.get("total_value")),
                _f((order.get("requested_by_profile") or {}).get("full_name")
                   or (order.get("requested_by_profile") or {}).get("email")),
            ])
    return workbook_bytes(wb)


def order_attachment_workbook(order: dict, items: list[dict]) -> bytes:
    """A single-order Excel file for the manager email attachment."""
    wb = Workbook()
    ws = _sheet(wb, "Order")
    store = order.get("stores") or {}
    supplier = order.get("suppliers") or {}
    requester = order.get("requested_by_profile") or {}

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 40

    ws.append(["Order ID", _f(order.get("po_number"))])
    ws.append(["Store", _f(store.get("name"))])
    ws.append(["Store Code", _f(store.get("code"))])
    ws.append(["Store Address", _f(order.get("delivery_address")
                                   or store.get("address") or store.get("location"))])
    ws.append(["Store Manager", _f(requester.get("full_name") or requester.get("email"))])
    ws.append(["Order Date", _f(order.get("requested_date") or order.get("created_at"))])
    ws.append(["Expected Date", _f(order.get("expected_date"))])
    ws.append(["Status", _f(order.get("status"))])
    ws.append(["Warehouse Status", _f(order.get("fulfillment_status"))])
    ws.append(["Priority", _f(order.get("priority"))])
    ws.append(["Supplier", _f(supplier.get("name"))])
    ws.append(["Notes", _f(order.get("notes"))])
    ws.append([])

    header_row = 14
    headers = ["Product Name", "SKU", "MRP", "Quantity", "Unit Price", "Line Total",
               "Product Image URL"]
    widths = [40, 20, 12, 12, 14, 16, 46]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        if widths[col - 1]:
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.freeze_panes = f"A{header_row + 1}"

    row_number = header_row
    for item in items:
        product = item.get("products") or {}
        row_number += 1
        values = [
            _f(product.get("name")),
            _f(product.get("sku")),
            _f(product.get("mrp")),
            _f(item.get("quantity")),
            _f(item.get("unit_price")),
            _f(item.get("line_total")),
            _f(product.get("image_url")),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_number, column=col, value=value)

    ws.cell(row=row_number + 2, column=1, value="Total Units").font = Font(bold=True)
    ws.cell(row=row_number + 2, column=4,
            value=sum(float(i.get("quantity") or 0) for i in items)).font = Font(bold=True)
    ws.cell(row=row_number + 3, column=1, value="Total Value").font = Font(bold=True)
    ws.cell(row=row_number + 3, column=6,
            value=_f(order.get("total_value"))).font = Font(bold=True)

    return workbook_bytes(wb)


# ---------------------------------------------------------------- packing list


def packing_list_workbook(packing: dict, items: list[dict]) -> bytes:
    """The warehouse packing list for one store order.

    Built from the packing_lists / packing_list_items rows that were saved when
    the list was generated, so re-downloading it months later reproduces exactly
    what was packed rather than recomputing from the order.
    """
    wb = Workbook()
    ws = _sheet(wb, "Packing List")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 44

    ws.cell(row=1, column=1, value="PACKING LIST").font = Font(bold=True, size=14,
                                                               color="7A1F2B")
    ws.append([])
    ws.append(["Packing List No.", _f(packing.get("packing_number"))])
    ws.append(["Order ID", _f(packing.get("po_number"))])
    ws.append(["Store", _f(packing.get("store_name"))])
    ws.append(["Store Address", _f(packing.get("store_address"))])
    ws.append(["Store Manager", _f(packing.get("store_manager"))])
    ws.append(["Warehouse", _f(packing.get("warehouse_name"))])
    ws.append(["Warehouse Staff", _f(packing.get("packed_by_name"))])
    ws.append(["Packing Date", _f(packing.get("packing_date"))])
    ws.append(["Total Lines", _f(packing.get("total_lines"))])
    ws.append(["Total Units", _f(packing.get("total_units"))])
    ws.append(["Notes", _f(packing.get("notes"))])
    ws.append([])

    header_row = 15
    headers = ["#", "Product Name", "SKU", "MRP", "Quantity Ordered",
               "Quantity Packed", "Product Image URL", "Packed ✓"]
    widths = [6, 42, 20, 12, 18, 18, 46, 12]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        if widths[col - 1]:
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.freeze_panes = f"A{header_row + 1}"

    row_number = header_row
    for index, item in enumerate(items, start=1):
        row_number += 1
        values = [
            index,
            _f(item.get("product_name")),
            _f(item.get("sku")),
            _f(item.get("mrp")),
            _f(item.get("quantity")),
            _f(item.get("packed_quantity")),
            _f(item.get("image_url")),
            "",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_number, column=col, value=value)

    ws.cell(row=row_number + 2, column=2, value="TOTAL").font = Font(bold=True)
    ws.cell(row=row_number + 2, column=5,
            value=sum(float(i.get("quantity") or 0) for i in items)).font = Font(bold=True)
    ws.cell(row=row_number + 2, column=6,
            value=sum(float(i.get("packed_quantity") or 0) for i in items)).font = Font(bold=True)

    ws.cell(row=row_number + 5, column=1, value="Packed by: ______________________")
    ws.cell(row=row_number + 5, column=4, value="Checked by: ______________________")
    ws.cell(row=row_number + 7, column=1, value="Received by (store): ______________________")
    ws.cell(row=row_number + 7, column=4, value="Date: ______________")

    return workbook_bytes(wb)


# ---------------------------------------------------------------- forecast


def forecast_workbook(rows: list[dict]) -> bytes:
    """Forecast results. `rows` come from the forecast export endpoint."""
    wb = Workbook()
    ws = _sheet(wb, "Forecast")
    headers = [
        "Product", "SKU", "Store", "Forecast Period", "Historical Demand Summary",
        "Historical Avg / Day", "Forecasted Demand", "Trend", "Trend %",
        "Recommended Action", "Recommendation Priority", "Basis", "Confidence",
    ]
    widths = [34, 18, 24, 26, 40, 16, 18, 12, 12, 46, 18, 22, 12]
    _style_sheet(ws, headers, widths)

    for row in rows:
        ws.append([
            _f(row.get("name")),
            _f(row.get("sku")),
            _f(row.get("store_name")),
            _f(row.get("period")),
            _f(row.get("history_summary")),
            _f(row.get("avg_daily_demand")),
            _f(row.get("forecasted_demand")),
            _f(row.get("trend")),
            _f(row.get("trend_pct")),
            _f(row.get("recommended_action")),
            _f(row.get("priority")),
            _f(row.get("basis_label")),
            _f(row.get("confidence")),
        ])
    return workbook_bytes(wb)